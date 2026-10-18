"""Synthetic fixtures only. Profiler runs as a standalone script in production's sandbox."""
import csv
import io
import json
from datetime import date

import pytest

from app.runner.profiler import alias_input, fingerprint, identity, profile
from app.services.dataset_profiles import apply_aliases


def test_csv_sample_contains_no_original_cells_and_is_reproducible(tmp_path):
    path = tmp_path / "data.csv"
    rows = [["customer", "amount", "day", "group", "free_text"]] + [
        [f"Private-{i}", f"{125.25 + i}", "2026-01-15", "rare-value" if i == 0 else "common-value",
         f"Confidential-{i}"] for i in range(30)]
    with path.open("w", newline="") as f:
        csv.writer(f).writerows(rows)
    schema, sample, sha = profile(path, "DS-0001")
    assert (schema, sample, sha) == profile(path, "DS-0001")
    cells = {c for r in rows[1:] for c in r}
    generated = list(csv.reader(io.StringIO(sample)))
    assert len(generated) == 21
    assert not cells.intersection(c for r in generated[1:] for c in r if c)
    assert not {identity(c) for r in rows[1:] for c in r}.intersection(
        identity(c) for r in generated[1:] for c in r if c)
    assert not any(v in sample for v in ("Private-", "Confidential-", "common-value", "rare-value"))
    cols = schema["sheets"][0]["columns"]
    assert cols[0]["role"] == "identifier" and "values" not in cols[0]
    assert cols[3]["values"] == ["common-value", "(rare)"]
    assert len(sha) == 64


def test_fingerprint_ignores_stats_but_tracks_structure(tmp_path):
    path = tmp_path / "data.csv"
    path.write_text("id,amount\na,12.50\nb,12.50\n")
    schema, _, sha = profile(path, "DS-0001")
    schema["dataset"] = "DS-9999"
    sheet = schema["sheets"][0]
    sheet["rows"] = 500
    sheet["columns"][0].update(nulls=0.8, cardinality=1, role="category", values=["x"])
    assert fingerprint(schema) == sha
    sheet["columns"][0]["type"] = "integer"
    assert fingerprint(schema) != sha


def test_xlsx_multi_sheet_and_aliases(tmp_path, monkeypatch):
    from openpyxl import Workbook, load_workbook
    path = tmp_path / "data.xlsx"
    book = Workbook()
    for sheet in (book.active, book.create_sheet("Other")):
        sheet.append(["email", "amount", "day"])
        for _ in range(10):
            sheet.append(["confidential@example.invalid", -12.75, date(2026, 2, 15)])
    book.save(path)
    schema, sample, sha = profile(path, "DS-0002")
    assert len(schema["sheets"]) == 2 and len(list(csv.reader(io.StringIO(sample)))) == 42
    assert "confidential@" not in sample and "-12.75" not in sample and "2026-02-15" not in sample
    original, sample = apply_aliases(schema, sample, {"s0": "Safe", "s0c0": "contact"})
    assert fingerprint(schema) != sha and "contact" in sample
    monkeypatch.setenv("TBR_INPUT", str(tmp_path))
    alias_input({"original": original, "aliases": {"s0": "Safe", "s0c0": "contact"}})
    import os
    from pathlib import Path
    renamed = load_workbook(Path(os.environ["TBR_INPUT"]) / "data.xlsx")
    assert renamed.worksheets[0].title == "Safe" and renamed.worksheets[0].cell(1, 1).value == "contact"
    assert load_workbook(path).worksheets[0].cell(1, 1).value == "email"


def test_alias_csv_preserves_values_and_delimiter(tmp_path, monkeypatch):
    path = tmp_path / "data.csv"
    path.write_text("secret;amount\nprivate;12.50\n")
    schema, _, _ = profile(path, "DS-0001")
    monkeypatch.setenv("TBR_INPUT", str(tmp_path))
    alias_input({"original": schema, "aliases": {"s0c0": "contact"}})
    import os
    from pathlib import Path
    assert (Path(os.environ["TBR_INPUT"]) / "data.csv").read_text() == "contact;amount\nprivate;12.50\n"


@pytest.mark.parametrize("body", ["a,a\n1,2\n", "a\n1,2\n"])
def test_ambiguous_columns_fail_closed(tmp_path, body):
    path = tmp_path / "data.csv"
    path.write_text(body)
    with pytest.raises(ValueError):
        profile(path, "DS-0001")


def test_aliases_must_be_unique(tmp_path):
    from app.errors import Unprocessable
    path = tmp_path / "data.csv"
    path.write_text("a,b\n1,2\n")
    schema, sample, _ = profile(path, "DS-0001")
    with pytest.raises(Unprocessable):
        apply_aliases(json.loads(json.dumps(schema)), sample, {"s0c0": "b"})


def test_excel_expansion_limit_before_parsing(tmp_path, monkeypatch):
    from zipfile import ZipFile
    from app.runner import profiler
    path = tmp_path / "data.xlsx"
    with ZipFile(path, "w") as z:
        z.writestr("oversized.xml", b"x" * 1024)
    monkeypatch.setattr(profiler, "MAX_UNPACKED", 512)
    with pytest.raises(ValueError, match="expansion limit"):
        profile(path, "DS-0001")


def test_excel_xml_entities_are_refused(tmp_path):
    from zipfile import ZipFile
    from openpyxl import Workbook
    path = tmp_path / "data.xlsx"
    Workbook().save(path)
    with ZipFile(path) as z:
        entries = {name: z.read(name) for name in z.namelist()}
    entries["xl/workbook.xml"] = b'''<?xml version="1.0"?>
<!DOCTYPE workbook [<!ENTITY private "private-source-cell">]>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">&private;</workbook>'''
    with ZipFile(path, "w") as z:
        for name, data in entries.items():
            z.writestr(name, data)
    from defusedxml.common import DefusedXmlException
    with pytest.raises(ValueError) as error:
        profile(path, "DS-0001")
    assert isinstance(error.value.__cause__, DefusedXmlException)
